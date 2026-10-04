from io import BytesIO
from struct import pack_into
from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from pdf_hyperlink_extractor import ExtractionError, extract_docx_links
from tests.document_builders import add_hyperlink, save_docx


def test_docx_finds_tables_nested_tables_headers_and_footers():
    document = Document()
    document.add_paragraph("https://body.example")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "https://table.example"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "ftp://nested.example/file"
    section = document.sections[0]
    section.header.paragraphs[0].text = "https://header.example"
    add_hyperlink(section.footer.paragraphs[0], "mailto:footer@example.com", "Contact")
    section.first_page_header.paragraphs[0].text = "https://first-header.example"
    assert extract_docx_links(save_docx(document)) == [
        "https://body.example",
        "https://table.example",
        "ftp://nested.example/file",
        "https://header.example",
        "mailto:footer@example.com",
        "https://first-header.example",
    ]


def test_docx_preserves_run_continuity_and_link_order():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Visit https://exa")
    paragraph.add_run("mple.com?q=a,b. ").bold = True
    add_hyperlink(paragraph, "https://target.example/path.", "https://label.example")
    paragraph.add_run(" mailto:a@example.com")
    assert extract_docx_links(save_docx(document)) == [
        "https://example.com?q=a,b",
        "https://target.example/path.",
        "https://label.example",
        "mailto:a@example.com",
    ]


def test_docx_keeps_paragraphs_and_line_breaks_separate():
    document = Document()
    document.add_paragraph("https://first.example")
    document.add_paragraph("https://second.example")
    paragraph = document.add_paragraph("https://third.example")
    paragraph.add_run().add_break()
    paragraph.add_run("https://fourth.example")
    paragraph.add_run().add_tab()
    paragraph.add_run("https://fifth.example")
    assert extract_docx_links(save_docx(document)) == [
        f"https://{name}.example" for name in ["first", "second", "third", "fourth", "fifth"]
    ]


def test_docx_text_boxes_and_content_controls():
    document = Document()
    paragraph = document.add_paragraph("https://outside.example ")
    paragraph._p.append(
        parse_xml(
            f"<w:r {nsdecls('w')}><w:pict><w:txbxContent>"
            "<w:p><w:r><w:t>https://textbox.example</w:t></w:r></w:p>"
            "</w:txbxContent></w:pict></w:r>"
        )
    )
    document._element.body.append(
        parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtContent>"
            "<w:p><w:r><w:t>https://control.example</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt>"
        )
    )
    assert extract_docx_links(save_docx(document)) == [
        "https://outside.example",
        "https://textbox.example",
        "https://control.example",
    ]


@pytest.mark.parametrize(
    "kind, content_type, rel_type",
    [
        ("footnotes", CT.WML_FOOTNOTES, RT.FOOTNOTES),
        ("endnotes", CT.WML_ENDNOTES, RT.ENDNOTES),
        ("comments", CT.WML_COMMENTS, RT.COMMENTS),
    ],
)
def test_docx_note_parts(kind, content_type, rel_type):
    document = Document()
    xml = parse_xml(
        f"<w:{kind} {nsdecls('w')}><w:p><w:r>"
        f"<w:t>https://{kind}.example</w:t></w:r></w:p></w:{kind}>"
    )
    part = XmlPart(PackURI(f"/word/{kind}.xml"), content_type, xml, document.part.package)
    document.part.relate_to(part, rel_type)
    assert extract_docx_links(save_docx(document)) == [f"https://{kind}.example"]


def test_docx_simple_and_complex_hyperlink_fields():
    document = Document()
    paragraph = document.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'HYPERLINK "https://simple.example" \\l "section"')
    paragraph._p.append(field)
    paragraph = document.add_paragraph()
    for tag, value in [
        ("w:fldChar", "begin"),
        ("w:instrText", ' HYPERLINK "https://complex.'),
        ("w:instrText", 'example" '),
        ("w:fldChar", "separate"),
        ("w:t", "Display label"),
        ("w:fldChar", "end"),
    ]:
        run = OxmlElement("w:r")
        element = OxmlElement(tag)
        if tag == "w:fldChar":
            element.set(qn("w:fldCharType"), value)
        else:
            element.text = value
        run.append(element)
        paragraph._p.append(run)
    assert extract_docx_links(save_docx(document)) == [
        "https://simple.example#section",
        "https://complex.example",
    ]


def test_docx_external_anchor_and_internal_links():
    document = Document()
    paragraph = document.add_paragraph()
    add_hyperlink(paragraph, "https://example.com/path", "Link")
    paragraph._p[-1].set(qn("w:anchor"), "section")
    internal = OxmlElement("w:hyperlink")
    internal.set(qn("w:anchor"), "internal-bookmark")
    paragraph._p.append(internal)
    assert extract_docx_links(save_docx(document)) == ["https://example.com/path#section"]


def test_docx_does_not_report_unused_relationships_or_internal_fields():
    document = Document()
    document.part.relate_to("https://unused.example", RT.HYPERLINK, is_external=True)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'HYPERLINK \\l "internal-bookmark"')
    document.add_paragraph()._p.append(field)
    assert extract_docx_links(save_docx(document)) == []


@pytest.mark.parametrize("data", [b"", b"not a docx", b"PK\x03\x04invalid"])
def test_docx_bad_input(data):
    with pytest.raises(ExtractionError, match="Unable to read DOCX"):
        extract_docx_links(data)


def test_docx_bad_xml_is_reported():
    data = save_docx(Document())
    buffer = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(buffer, "w") as corrupt:
        for item in original.infolist():
            corrupt.writestr(
                item,
                b"<broken"
                if item.filename == "word/document.xml"
                else original.read(item.filename),
            )
    with pytest.raises(ExtractionError, match="Unable to read DOCX"):
        extract_docx_links(buffer)


@pytest.mark.parametrize(
    "instruction, expected",
    [
        ('HYPERLINK \\l "intro" "https://example.com" \\o "Tooltip"', "https://example.com#intro"),
        (r"HYPERLINK C:\docs\file.docx \\* MERGEFORMAT", r"C:\docs\file.docx"),
        (r"HYPERLINK \\server\share\file.docx", r"\\server\share\file.docx"),
        ("HYPERLINK ftp://example.com/file \\n", "ftp://example.com/file"),
        ('HYPERLINK "https://example.com#existing" \\l "other"', "https://example.com#existing"),
        ('HYPERLINK ""', None),
        ("PAGE", None),
    ],
)
def test_docx_hyperlink_field_switches_and_file_targets(instruction, expected):
    document = Document()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    document.add_paragraph()._p.append(field)
    assert extract_docx_links(save_docx(document)) == ([] if expected is None else [expected])


def test_docx_empty_hyperlink_target_still_reads_display_text():
    document = Document()
    add_hyperlink(document.add_paragraph(), "", "https://label.example")
    assert extract_docx_links(save_docx(document)) == ["https://label.example"]


@pytest.mark.parametrize(
    "local_offset, central_offset, value",
    [
        pytest.param(8, 10, 99, id="unsupported-compression"),
        pytest.param(6, 8, 1, id="encrypted-entry"),
    ],
)
def test_unreadable_zip_entries_are_friendly_errors(local_offset, central_offset, value):
    data = bytearray(save_docx(Document()))
    with ZipFile(BytesIO(data)) as archive:
        local_header = archive.infolist()[0].header_offset
        central_directory = archive.start_dir
    pack_into("<H", data, local_header + local_offset, value)
    pack_into("<H", data, central_directory + central_offset, value)
    with pytest.raises(ExtractionError, match="Unable to read DOCX"):
        extract_docx_links(bytes(data))
