import sys
from io import BytesIO
from pathlib import Path

import pytest
from PyPDF2 import PdfReader, PdfWriter
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.append(str(Path(__file__).resolve().parents[1]))

from pdf_extractor import extract_pdf_links
from streamlit_app import extract_docx_links


PDF_HEADER = b"%PDF-1.4\n"
PDF_STREAM_TEXT = "BT /F1 12 Tf 72 720 Td (Visit https://example.org) Tj ET"


def build_sample_pdf_bytes() -> bytes:
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R /Annots [6 0 R] >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(PDF_STREAM_TEXT)} >>\nstream\n{PDF_STREAM_TEXT}\nendstream",
        "<< /Type /Annot /Subtype /Link /Rect [72 700 200 720] /Border [0 0 0] /A << /S /URI /URI (https://example.com) >> >>",
    ]

    body = bytearray()
    offsets = [0]
    current = len(PDF_HEADER)
    for index, obj in enumerate(objects, start=1):
        obj_bytes = f"{index} 0 obj\n{obj}\nendobj\n".encode("utf-8")
        offsets.append(current)
        body.extend(obj_bytes)
        current += len(obj_bytes)

    xref_offset = len(PDF_HEADER) + len(body)
    lines = ["xref", f"0 {len(objects) + 1}", "0000000000 65535 f "]
    for offset in offsets[1:]:
        lines.append(f"{offset:010} 00000 n ")
    xref = "\n".join(lines).encode("ascii") + b"\n"
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return PDF_HEADER + bytes(body) + xref + trailer


def build_encrypted_pdf_bytes() -> bytes:
    reader = PdfReader(BytesIO(build_sample_pdf_bytes()))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("secret")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_sample_docx_bytes() -> bytes:
    document = Document()
    paragraph = document.add_paragraph("Example link: ")
    add_hyperlink(paragraph, "https://example.com", "Example")
    document.add_paragraph("Another https://example.org reference")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_pdf_links():
    pdf_file = BytesIO(build_sample_pdf_bytes())
    links = extract_pdf_links(pdf_file)
    assert set(links) == {"https://example.com", "https://example.org"}


def test_extract_pdf_links_encrypted():
    encrypted_pdf = BytesIO(build_encrypted_pdf_bytes())
    with pytest.raises(ValueError):
        extract_pdf_links(encrypted_pdf)


def test_extract_docx_links():
    docx_file = BytesIO(build_sample_docx_bytes())
    links = extract_docx_links(docx_file)
    assert set(links) == {"https://example.com", "https://example.org"}
