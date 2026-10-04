"""Generate real documents in memory for extraction and UI regression tests."""

from io import BytesIO

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfWriter
from pypdf.annotations import Link
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def build_sample_pdf_bytes() -> bytes:
    return build_pdf("Visit https://example.org", ("https://example.com",))


def build_encrypted_pdf_bytes(password: str = "secret") -> bytes:
    return build_pdf("Visit https://example.org", ("https://example.com",), password=password)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_sample_docx_bytes() -> bytes:
    document = Document()
    paragraph = document.add_paragraph("Example link: ")
    add_hyperlink(paragraph, "https://example.com", "Example")
    document.add_paragraph("Another https://example.org reference")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def save_docx(document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(
    text: str = "",
    targets: tuple[str, ...] = (),
    *,
    page_count: int = 1,
    password: str | None = None,
    algorithm: str = "AES-256",
    indirect_actions: bool = False,
    base_uri: str | None = None,
) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for _ in range(page_count):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
            }
        )
        if text:
            instructions = ["BT /F1 12 Tf 72 720 Td"]
            for line in text.splitlines():
                escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                instructions.append(f"({escaped}) Tj 0 -14 Td")
            instructions.append("ET")
            content = DecodedStreamObject()
            content.set_data("\n".join(instructions).encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(content)
        for target in targets:
            annotation = Link(rect=(72, 700, 200, 720), url=target)
            if indirect_actions:
                annotation[NameObject("/A")] = writer._add_object(annotation["/A"])
            writer.add_annotation(page, annotation)
    if base_uri is not None:
        from pypdf.generic import TextStringObject

        writer.root_object[NameObject("/URI")] = DictionaryObject(
            {
                NameObject("/Base"): TextStringObject(base_uri),
            }
        )
    if password is not None:
        writer.encrypt(password, algorithm=algorithm)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
