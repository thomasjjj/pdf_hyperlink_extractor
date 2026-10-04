"""Measure extraction on generated large documents without timing generation."""

import argparse
from collections.abc import Callable
from statistics import median
from time import perf_counter

from docx import Document

from extractors import extract_docx_links, extract_pdf_links
from tests.document_builders import build_pdf, save_docx


def _positive_int(value: str) -> int:
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("must be a positive integer")
    return number


def _measure(extractor: Callable[[bytes], list[str]], data: bytes, runs: int) -> float:
    samples = []
    for _ in range(runs):
        start = perf_counter()
        extractor(data)
        samples.append(perf_counter() - start)
    return median(samples)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--paragraphs", type=_positive_int, default=5000)
    parser.add_argument("--pages", type=_positive_int, default=100)
    parser.add_argument("--runs", type=_positive_int, default=5)
    options = parser.parse_args()

    document = Document()
    for index in range(options.paragraphs):
        document.add_paragraph(f"Reference https://example.com/item/{index % 100}")
    docx = save_docx(document)
    expected = [
        f"https://example.com/item/{index}" for index in range(min(options.paragraphs, 100))
    ]
    if extract_docx_links(docx) != expected:
        raise RuntimeError("DOCX benchmark returned incorrect links")

    pdf = build_pdf(
        "https://example.com mailto:a@example.com ftp://example.com/file",
        ("https://example.com",),
        page_count=options.pages,
    )
    if extract_pdf_links(pdf) != [
        "https://example.com",
        "mailto:a@example.com",
        "ftp://example.com/file",
    ]:
        raise RuntimeError("PDF benchmark returned incorrect links")

    print(f"Median of {options.runs} runs, excluding fixture generation:")
    docx_seconds = _measure(extract_docx_links, docx, options.runs)
    pdf_seconds = _measure(extract_pdf_links, pdf, options.runs)
    print(f"DOCX ({options.paragraphs} paragraphs): {docx_seconds:.4f}s")
    print(f"PDF ({options.pages} pages): {pdf_seconds:.4f}s")


if __name__ == "__main__":
    main()
